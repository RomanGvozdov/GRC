"""Збирає багатий каталог НД ТЗІ 3.6-006-24 (повний 800-53 Rev 5, укр.) у seed-JSON.

Джерела (у docs/nd-tzi/):
- nist_800-53r5_nd-tzi_controls.xlsx — аркуш «НД ТЗІ 3.6-006-24 (UA)»: 1189 заходів
  (текст, рекомендації, пов'язані, посилання);
- galuzevyi_profil_reestriv.docx — галузевий профіль реєстрів (≈119 заходів);
- app/seed_data/nd_tzi_3_6_006_24.json — членство конфіденц./службова (84/97).

Вихід: app/seed_data/nd_tzi_catalog.json — єдиний каталог з членством у профілях
(confidential/service/registry); baselines будуються з нього при сидингу.

Запуск: cd backend && .venv/bin/python scripts/build_catalogs.py
"""
import json
import re
from pathlib import Path

import docx
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT.parent / "docs" / "nd-tzi"
XLSX = DOCS / "nist_800-53r5_nd-tzi_controls.xlsx"
REGISTRY_DOCX = DOCS / "galuzevyi_profil_reestriv.docx"
CONF_DOCX = DOCS / "Базовий_профіль_для_конфіденційної_та_відкритої (1).docx"
SERV_DOCX = DOCS / "Базовий_профіль_безпеки_службова_інформація (1).docx"
# Порядок переваги для витягу ODP-параметрів (найповніший профіль — перший)
PROFILE_DOCX = [REGISTRY_DOCX, SERV_DOCX, CONF_DOCX]
PROFILE_FILE = ROOT / "app" / "seed_data" / "nd_tzi_3_6_006_24.json"
OUT = ROOT / "app" / "seed_data" / "nd_tzi_catalog.json"
OUT_EN = ROOT / "app" / "seed_data" / "nist_800_53_r5_en.json"

CATALOG_CODE = "nd-tzi-3-6-006-24"
CATALOG_NAME = "НД ТЗІ 3.6-006-24 (повний каталог, NIST 800-53 Rev 5)"
CATALOG_VERSION = "2024.r5.1"  # +ODP-параметри (org-defined + прописані значення)

EN_CATALOG_CODE = "nist-800-53-r5-en"
EN_CATALOG_NAME = "NIST SP 800-53 Rev. 5 (English)"
EN_CATALOG_VERSION = "Rev. 5"
EN_SHEET = "SP 800-53 Rev 5 (EN)"

_CYR = "АВСЕНІКМОРТХУ"
_LAT = "ABCEHIKMOPTXY"
_TR = str.maketrans(_CYR, _LAT)


def norm_code(raw: str) -> str | None:
    s = (raw or "").strip().translate(_TR)
    m = re.match(r"([A-Z]{2}-\d+(?:\(\d+\))?)", s)
    return m.group(1) if m else None


def registry_codes() -> set[str]:
    d = docx.Document(str(REGISTRY_DOCX))
    codes = set()
    for row in d.tables[0].rows:
        cells = [c.text for c in row.cells]
        code = norm_code(cells[2]) if len(cells) > 2 else None
        if code:
            codes.add(code)
    return codes


_ODP_GENERIC = re.compile(r"Призначенн|Вибір|Надання|визначен", re.IGNORECASE)


def _profile_text_by_code(path: Path) -> dict[str, str]:
    d = docx.Document(str(path))
    out: dict[str, str] = {}
    for row in d.tables[0].rows:
        cells = [c.text for c in row.cells]
        if len(cells) < 3:
            continue
        code = norm_code(cells[2])
        if code:
            out[code] = cells[2].strip()
    return out


def _params_from_text(code: str, text: str) -> list[dict]:
    """ODP-параметри контролю з тексту базового профілю.

    org_defined (`[Призначення/Вибір/Надання: …]`) — заповнює організація (без default);
    прописане значення (`[24 годин]`, `[≤30 хвилин]`) — default_value (мінімум за НД ТЗІ),
    підпис береться з контексту перед дужкою."""
    cid = code.lower().replace("(", ".").replace(")", "")
    params, seen, n = [], set(), 0
    for m in re.finditer(r"\[[^\[\]]+\]", text):
        body = m.group(0)[1:-1].strip()
        if not body or body in seen:
            continue
        seen.add(body)
        n += 1
        org = bool(_ODP_GENERIC.search(body))
        if org:
            label, default = body[:300], None
        else:
            ctx = re.split(r"[;.\n]", text[: m.start()])[-1].strip()
            label = (ctx[-80:].strip() or "Прописане значення за НД ТЗІ")
            default = body[:300]
        params.append({
            "key": f"{cid}_odp.{n:02d}",
            "label": label,
            "default_value": default,
            "org_defined": org,
        })
    return params


def parameters_by_code() -> dict[str, list[dict]]:
    """ODP-параметри для всіх контролів базових профілів (об'єднання; найповніший
    профіль має перевагу за текстом)."""
    texts: dict[str, str] = {}
    for path in reversed(PROFILE_DOCX):  # conf → serv → registry: пізніший перезаписує
        if path.exists():
            texts.update(_profile_text_by_code(path))
    return {code: _params_from_text(code, t) for code, t in texts.items()}


def profile_membership() -> tuple[set[str], set[str]]:
    # Джерело членства конфіденц./службова: первинний профільний файл, а якщо його
    # вже прибрано — раніше згенерований каталог (self-contained, відтворюється).
    src = PROFILE_FILE if PROFILE_FILE.exists() else OUT
    data = json.loads(src.read_text(encoding="utf-8"))
    conf, serv = set(), set()
    for r in data["requirements"]:
        profiles = r.get("profiles") or []
        if "confidential" in profiles:
            conf.add(r["code"])
        if "service" in profiles:
            serv.add(r["code"])
    return conf, serv


def build_description(text: str, guidance: str) -> str:
    text = (text or "").strip()
    guidance = (guidance or "").strip()
    if guidance and guidance.lower() != "немає.":
        return f"{text}\n\nРекомендації з реалізації:\n{guidance}"
    return text


def build_english_catalog() -> int:
    """Англійський каталог NIST 800-53 Rev 5 (data-only, без профілів) — для OSCAL/
    NIST крос-референсу та майбутньої EN-локалізації."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[EN_SHEET]
    requirements, seen = [], set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        if code in seen:
            continue
        seen.add(code)
        description = str(row[2] or "").strip()
        discussion = str(row[3] or "").strip()
        if discussion and discussion.lower() != "none.":
            description = f"{description}\n\nDiscussion:\n{discussion}"
        requirements.append({
            "code": code,
            "title": str(row[1] or "").strip()[:500],
            "description": description,
        })
    out = {
        "code": EN_CATALOG_CODE,
        "name": EN_CATALOG_NAME,
        "version": EN_CATALOG_VERSION,
        "is_custom": False,
        "requirements": requirements,
    }
    OUT_EN.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    return len(requirements)


def main() -> None:
    conf, serv = profile_membership()
    registry = registry_codes()
    params_map = parameters_by_code()

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb["НД ТЗІ 3.6-006-24 (UA)"]

    requirements = []
    seen = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        code = str(row[0]).strip()
        if code in seen:
            continue
        seen.add(code)
        title = str(row[1] or "").strip()[:500]
        description = build_description(str(row[2] or ""), str(row[3] or ""))
        profiles = []
        if code in conf:
            profiles.append("confidential")
        if code in serv:
            profiles.append("service")
        if code in registry:
            profiles.append("registry")
        req = {
            "code": code,
            "title": title,
            "description": description,
            "profiles": profiles,
        }
        if params_map.get(code):
            req["parameters"] = params_map[code]
        requirements.append(req)

    out = {
        "code": CATALOG_CODE,
        "name": CATALOG_NAME,
        "version": CATALOG_VERSION,
        "is_custom": False,
        "requirements": requirements,
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"Записано {len(requirements)} заходів у {OUT.name}")
    print(f"  конфіденц.: {sum(1 for r in requirements if 'confidential' in r['profiles'])}")
    print(f"  службова:   {sum(1 for r in requirements if 'service' in r['profiles'])}")
    print(f"  реєстри:    {sum(1 for r in requirements if 'registry' in r['profiles'])}")
    miss = registry - seen
    if miss:
        print(f"  УВАГА: реєстрові коди поза каталогом: {sorted(miss)}")
    n_par = sum(len(r.get("parameters", [])) for r in requirements)
    n_pre = sum(1 for r in requirements for p in r.get("parameters", [])
                if not p["org_defined"])
    n_ctl = sum(1 for r in requirements if r.get("parameters"))
    print(f"  ODP-параметрів: {n_par} на {n_ctl} контролях "
          f"(прописаних значень: {n_pre}, org-defined: {n_par - n_pre})")

    en_count = build_english_catalog()
    print(f"Записано {en_count} заходів у {OUT_EN.name} (англ. каталог)")


if __name__ == "__main__":
    main()
