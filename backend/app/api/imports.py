"""Імпорт реєстрів ризиків і контролів з Excel.

Формат — той самий, що в експорті (українські заголовки). Двокроковий процес:
dry_run=true повертає звіт перевірки без збереження; dry_run=false — імпортує
валідні рядки (рядки з помилками пропускаються).
"""

import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.controls import next_code as next_control_code
from app.api.risks import next_code as next_risk_code
from app.core.deps import require_permission
from app.database import get_db
from app.models import (
    Control,
    ControlImplementation,
    InformationSystem,
    Requirement,
    Risk,
    RiskCategory,
    User,
)
from app.schemas_phase4 import ImportReport, ImportRowError
from app.services.audit import log_action

router = APIRouter(prefix="/imports", tags=["imports"])

RISK_HEADERS = [
    "Назва", "Опис", "Категорія", "Статус", "Відповідальний (email)", "Системи",
    "Ймовірність (притаманна)", "Вплив (притаманний)",
    "Ймовірність (залишкова)", "Вплив (залишковий)",
    "Стратегія обробки", "Дата наступного перегляду",
    "Активи", "Джерело загрози", "Вразливість",
]
CONTROL_HEADERS = [
    "Назва", "Опис", "Тип", "Відповідальний (email)", "Вимоги (коди через кому)",
    "Системи", "Статус впровадження", "Обґрунтування (не застосовно)",
    "Дата наступної перевірки",
]

_RISK_STATUS = {
    "чернетка": "draft", "ідентифікований": "identified", "оцінений": "assessed",
    "в обробці": "in_treatment", "моніториться": "monitored", "закритий": "closed",
}
_STRATEGY = {"зменшити": "mitigate", "прийняти": "accept", "уникнути": "avoid",
             "передати": "transfer"}
_IMPL = {"не впроваджено": "not_implemented", "частково": "partial",
         "впроваджено": "implemented", "не застосовно": "not_applicable"}
_CONTROL_TYPE = {"превентивний": "preventive", "детективний": "detective",
                 "коригувальний": "corrective"}


def _template(headers: list[str], example: list, filename: str) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    ws.append(example)
    out = io.BytesIO()
    wb.save(out)
    return StreamingResponse(
        io.BytesIO(out.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/template/risks")
def risks_template(_: User = Depends(require_permission("risks", "manage"))):
    return _template(
        RISK_HEADERS,
        ["Витік даних через фішинг", "Опис ризику", "Інформаційна безпека",
         "Ідентифікований", "user@company.ua", "Корпоративна мережа, Вебпортал",
         4, 5, 2, 3, "Зменшити", "2026-12-31", "Поштова система", "Зловмисник", "Фішинг"],
        "import_risks_template.xlsx",
    )


@router.get("/template/controls")
def controls_template(_: User = Depends(require_permission("controls", "manage"))):
    return _template(
        CONTROL_HEADERS,
        ["Резервне копіювання", "Щоденні бекапи з перевіркою відновлення",
         "Превентивний", "user@company.ua", "A.8.13, CP-9",
         "Корпоративна мережа", "Впроваджено", "", "2026-06-30"],
        "import_controls_template.xlsx",
    )


def _read_rows(file: UploadFile, expected_headers: list[str]):
    try:
        wb = load_workbook(io.BytesIO(file.file.read()), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Не вдалося прочитати файл .xlsx")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Файл порожній")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    if headers[: len(expected_headers)] != expected_headers:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Заголовки не збігаються з шаблоном. Завантажте актуальний шаблон імпорту",
        )
    return rows[1:]


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"невірна дата «{text}» (очікується РРРР-ММ-ДД або ДД.ММ.РРРР)")


def _parse_scale(value, label: str) -> int | None:
    if value is None or value == "":
        return None
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label}: «{value}» не є числом")
    if not 1 <= number <= 5:
        raise ValueError(f"{label}: {number} поза шкалою 1–5")
    return number


def _lookup(mapping: dict[str, str], value, label: str) -> str | None:
    text = _text(value)
    if text is None:
        return None
    key = text.lower()
    if key not in mapping:
        raise ValueError(f"{label}: невідоме значення «{text}»")
    return mapping[key]


def _get_or_create_systems(db: Session, value, cache: dict) -> list[InformationSystem]:
    text = _text(value)
    if not text:
        return []
    systems = []
    for name in [n.strip() for n in text.split(",") if n.strip()]:
        key = name.lower()
        if key not in cache:
            system = db.scalar(
                select(InformationSystem).where(InformationSystem.name == name)
            )
            if system is None:
                from app.api.systems import _next_code

                system = InformationSystem(code=_next_code(db), name=name)
                db.add(system)
                db.flush()
            cache[key] = system
        systems.append(cache[key])
    return systems


def _find_user(db: Session, value) -> int | None:
    email = _text(value)
    if not email:
        return None
    user = db.scalar(select(User).where(User.email == email.lower()))
    if user is None:
        raise ValueError(f"користувача з email «{email}» не знайдено")
    return user.id


@router.post("/risks", response_model=ImportReport)
def import_risks(
    dry_run: bool = True,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    actor: User = Depends(require_permission("risks", "manage")),
):
    rows = _read_rows(file, RISK_HEADERS)
    errors: list[ImportRowError] = []
    valid: list[dict] = []
    category_cache: dict = {}
    system_cache: dict = {}

    for index, row in enumerate(rows, start=2):
        row = list(row) + [None] * (len(RISK_HEADERS) - len(row))
        if not _text(row[0]):
            continue  # порожній рядок
        try:
            category = None
            if _text(row[2]):
                name = _text(row[2])
                key = name.lower()
                if key not in category_cache:
                    found = db.scalar(select(RiskCategory).where(RiskCategory.name == name))
                    if found is None:
                        found = RiskCategory(name=name)
                        db.add(found)
                        db.flush()
                    category_cache[key] = found
                category = category_cache[key]
            valid.append({
                "title": _text(row[0]),
                "description": _text(row[1]),
                "category_id": category.id if category else None,
                "status": _lookup(_RISK_STATUS, row[3], "Статус") or "identified",
                "owner_id": _find_user(db, row[4]),
                "systems": _get_or_create_systems(db, row[5], system_cache),
                "inherent_likelihood": _parse_scale(row[6], "Ймовірність (притаманна)"),
                "inherent_impact": _parse_scale(row[7], "Вплив (притаманний)"),
                "residual_likelihood": _parse_scale(row[8], "Ймовірність (залишкова)"),
                "residual_impact": _parse_scale(row[9], "Вплив (залишковий)"),
                "treatment_strategy": _lookup(_STRATEGY, row[10], "Стратегія"),
                "next_review_date": _parse_date(row[11]),
                "assets": _text(row[12]),
                "threat_source": _text(row[13]),
                "vulnerability": _text(row[14]),
            })
        except ValueError as exc:
            errors.append(ImportRowError(row=index, message=str(exc)))

    created = 0
    if not dry_run:
        for data in valid:
            systems = data.pop("systems")
            risk = Risk(code=next_risk_code(db), **data)
            risk.systems = systems
            db.add(risk)
            db.flush()
            created += 1
        log_action(db, actor, "import", "risk", None, {"created": created})
        db.commit()
    else:
        db.rollback()

    return ImportReport(
        total_rows=len(valid) + len(errors), valid_rows=len(valid),
        errors=errors, created=created, dry_run=dry_run,
    )


@router.post("/controls", response_model=ImportReport)
def import_controls(
    dry_run: bool = True,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    actor: User = Depends(require_permission("controls", "manage")),
):
    rows = _read_rows(file, CONTROL_HEADERS)
    errors: list[ImportRowError] = []
    valid: list[dict] = []
    system_cache: dict = {}
    requirements_by_code: dict[str, list[Requirement]] = {}

    def find_requirements(value) -> list[Requirement]:
        text = _text(value)
        if not text:
            return []
        result = []
        for code in [c.strip() for c in text.split(",") if c.strip()]:
            if code not in requirements_by_code:
                requirements_by_code[code] = list(
                    db.scalars(select(Requirement).where(Requirement.code == code)).all()
                )
            found = requirements_by_code[code]
            if not found:
                raise ValueError(f"вимогу з кодом «{code}» не знайдено в каталогах")
            result.extend(found)
        return result

    for index, row in enumerate(rows, start=2):
        row = list(row) + [None] * (len(CONTROL_HEADERS) - len(row))
        if not _text(row[0]):
            continue
        try:
            impl_status = _lookup(_IMPL, row[6], "Статус впровадження") or "not_implemented"
            na_justification = _text(row[7])
            if impl_status == "not_applicable" and not na_justification:
                raise ValueError('статус «не застосовно» вимагає обґрунтування')
            valid.append({
                "name": _text(row[0]),
                "description": _text(row[1]),
                "control_type": _lookup(_CONTROL_TYPE, row[2], "Тип"),
                "owner_id": _find_user(db, row[3]),
                "requirements": find_requirements(row[4]),
                "systems": _get_or_create_systems(db, row[5], system_cache),
                "impl_status": impl_status,
                "na_justification": na_justification,
                "next_review_date": _parse_date(row[8]),
            })
        except ValueError as exc:
            errors.append(ImportRowError(row=index, message=str(exc)))

    created = 0
    if not dry_run:
        for data in valid:
            control = Control(
                code=next_control_code(db),
                name=data["name"],
                description=data["description"],
                control_type=data["control_type"],
                owner_id=data["owner_id"],
            )
            control.requirements = data["requirements"]
            targets = data["systems"] or [None]
            for system in targets:
                control.implementations.append(ControlImplementation(
                    system_id=system.id if system else None,
                    implementation_status=data["impl_status"],
                    na_justification=data["na_justification"],
                    next_review_date=data["next_review_date"],
                ))
            db.add(control)
            db.flush()
            created += 1
        log_action(db, actor, "import", "control", None, {"created": created})
        db.commit()
    else:
        db.rollback()

    return ImportReport(
        total_rows=len(valid) + len(errors), valid_rows=len(valid),
        errors=errors, created=created, dry_run=dry_run,
    )
